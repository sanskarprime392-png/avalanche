# -*- coding: utf-8 -*-
"""Supplementary material: the invariance checks behind Section 4.1."""
import os

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DATA = r"C:\Users\Sanskar\Documents\avalanche_data"
RES = os.path.join(DATA, "results")
OUT = r"C:\Users\Sanskar\Documents\avalanche\Avalanche_Supplementary.docx"
NAVY = RGBColor(0x0B, 0x3D, 0x5C)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def H(t, lvl=1):
    h = doc.add_heading(t, level=lvl)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h


def P(t, size=11):
    p = doc.add_paragraph()
    p.add_run(t).font.size = Pt(size)
    return p


def table(headers, rows, widths, caption):
    c = doc.add_paragraph()
    r = c.add_run(caption)
    r.bold = True
    r.font.size = Pt(9.5)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        rr = cell.paragraphs[0].add_run(h)
        rr.bold = True
        rr.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            cells[j].paragraphs[0].add_run(str(v)).font.size = Pt(9)
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
    doc.add_paragraph()


t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Supplementary material")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = NAVY
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Reported accuracy in geohazard susceptibility mapping is a function of the "
              "absence sampling design")
r.italic = True
r.font.size = Pt(10.5)
doc.add_paragraph()

H("S1  Invariance of the inflation to simulation settings", 1)
P("Section 4.1 reports that a model with no predictive skill inflates with absence displacement. "
  "A synthetic result invites the objection that it reflects the construction rather than the "
  "problem, so the null arm was repeated across six settings varying the learner, the number of "
  "predictors and the sample size. Null-model AUC differs by at most 0.030 at any displacement and "
  "by 0.001 at δ = 6.")
rb = pd.read_csv(os.path.join(RES, "simulation_robustness.csv"))
piv = rb.pivot(index="setting", columns="delta", values="auc")
table(["Setting"] + [f"δ={c:g}" for c in piv.columns],
      [[i] + [f"{piv.loc[i, c]:.3f}" for c in piv.columns] for i in piv.index],
      [2.3] + [0.62] * len(piv.columns),
      "Table S1. Null-model AUC under varying simulation settings.")

H("S2  Invariance of the inflation to displacement direction", 1)
P("The construction displaces absences along a single axis carrying no information about the "
  "outcome, which raises the question of whether that direction is special. Distributing the same "
  "displacement across five axes, along a random direction in the uninformative subspace, or "
  "partly along the informative axis itself — which is what an absence rule favouring gentle "
  "terrain does — changes null-arm AUC by at most 0.016 at any displacement. The inflation is a "
  "function of distance, not of direction.")
dt = pd.read_csv(os.path.join(RES, "direction_test.csv"))
nul = dt[dt.arm == "no skill"] if "arm" in dt.columns else dt
piv = nul.pivot(index="direction", columns="delta", values="auc")
table(["Displacement direction"] + [f"δ={c:g}" for c in piv.columns],
      [[i] + [f"{piv.loc[i, c]:.3f}" for c in piv.columns] for i in piv.index],
      [2.1] + [0.62] * len(piv.columns),
      "Table S2. Null-model AUC under different displacement directions.")

H("S3  Failure of subtraction as a correction", 1)
P("Section 4.2 withdraws the proposal to subtract a null baseline from a reported AUC. The "
  "evidence is below. Holding true skill fixed and varying only the displacement, the subtraction "
  "should return a constant; instead its error grows with displacement across every data structure "
  "tested, and fastest for clustered presences, which is what real inventories look like.")
cv = pd.read_csv(os.path.join(RES, "correction_validation.csv"))
piv = cv.pivot_table(index="structure", columns="delta", values="error",
                     aggfunc=lambda v: float(np.mean(np.abs(v))))
nice = {"gaussian": "Gaussian predictors", "correlated": "Correlated predictors",
        "nonlinear": "Nonlinear boundary", "clustered": "Clustered presences",
        "label_noise": "20% label noise"}
table(["Data structure"] + [f"δ={c:g}" for c in piv.columns],
      [[nice.get(i, i)] + [f"{piv.loc[i, c]:.3f}" for c in piv.columns] for i in piv.index],
      [2.1] + [0.75] * len(piv.columns),
      "Table S3. Mean absolute error of the subtraction, by data structure and displacement.")

th = pd.read_csv(os.path.join(RES, "correction_threshold.csv"))
g = th.groupby("delta").agg(h=("headroom", "mean"), a=("adjusted", "mean"),
                            r=("reference", "mean"), e=("error", lambda v: v.abs().mean()))
table(["δ (SD)", "Room left on the AUC scale", "Mean |error|", "True skill recovered"],
      [[f"{d:g}", f"{row.h:.3f}", f"{row.e:.3f}", f"{100 * row.a / row.r:.0f}%"]
       for d, row in g.iterrows()],
      [0.9, 2.0, 1.2, 1.6],
      "Table S4. Recovery of known true skill by subtraction, against absence displacement.")
P("The phrase reported as true skill recovered is the ratio of two quantities, both in AUC units: the "
  "numerator is the reported AUC at that displacement minus the null-model AUC at the same "
  "displacement, and the denominator is the reference skill, defined as the AUC obtained with "
  "absences drawn without displacement, minus 0.5. A working correction would return 100% at every "
  "displacement, because true skill is held constant by construction. The decay instead tracks the "
  "room remaining on the AUC scale, which is bounded above by unity: at δ = 3 that room is 0.020 "
  "and the subtraction recovers 4% of the known skill. The information is destroyed rather than "
  "displaced, which is why Section 4.2 reports an identifiability limit instead of a correction.")

H("S4  Sensitivity to the background slope window", 1)
P("The background slope window sets the separability of the absence design directly. Across four "
  "alternatives, with presences, predictors and model unchanged, reported AUC spans 0.031.")
try:
    sw = pd.read_csv(os.path.join(RES, "slope_window_sensitivity.csv"))
    table(["Slope window", "n", "AUC", "95% CI", "Separability"],
          [[f"{r.window}°", f"{int(r.n):,}", f"{r.auc:.3f}",
            f"[{r.ci_low:.3f}, {r.ci_high:.3f}]", f"{r.delta:.2f}"] for _, r in sw.iterrows()],
          [1.3, 0.8, 0.9, 1.5, 1.2],
          "Table S5. Effect of the background slope window.")
except FileNotFoundError:
    P("(slope_window_sensitivity.csv not found)")

H("S5  Applying the comparison to another study", 1)
P("designcheck.py in the repository runs the design comparison on any presence-only susceptibility "
  "dataset, independently of the rest of this code. Given a table of predictor values at event "
  "locations and a pool of candidate absence locations, it scores one model against several "
  "absence designs and reports the spread. Run on this study's own data with a pool spanning valley "
  "floor to avalanche-capable terrain, drawing absences at random from the pool — the common "
  "default — returns AUC 0.954 at a separability of 3.06, while the three matched designs return "
  "0.923 to 0.930 at separabilities of 2.09 to 2.15. The spread across all four designs is 0.031, "
  "and 0.024 between the random default and the most thoroughly matched design. Either figure is "
  "produced entirely by the sampling, with presences, predictors and model unchanged.")

H("S6  Invariance to the presence-to-background ratio", 1)
P("Susceptibility studies draw background points at ratios from 1:1 to 1:10 and the choice is "
  "rarely justified. Because a ratio change alters both the class balance and the effective sample "
  "size, the null arm was repeated across four ratios. Null-model AUC varies by at most 0.024 at "
  "any displacement, so the inflation reported in Section 4.1 is not an artefact of the 1:1 design "
  "used elsewhere in this work.")
pv = pd.read_csv(os.path.join(RES, "prevalence_test.csv"))
piv = pv.pivot(index="ratio", columns="delta", values="auc")
table(["Presence : background"] + [f"d={c:g}" for c in piv.columns],
      [[i] + [f"{piv.loc[i, c]:.3f}" for c in piv.columns] for i in piv.index],
      [1.8] + [0.62] * len(piv.columns),
      "Table S6. Null-model AUC by presence-to-background ratio.")

H("S7  Statistical reporting", 1)
P("All simulation figures are means over independent replications, each drawing a fresh dataset "
  "and refitting the model. Seeds are the replication index, beginning at 0, so every result is "
  "reproducible by re-running the corresponding function with its default arguments. Replication "
  "counts are 20 for the curve in Table 1, 8 for the invariance checks in Tables S1 and S2, 4 for "
  "the correction-validation grid in Table S3, 3 for the fine-grained sweep in Table S4 and 6 for "
  "the ratio check in Table S6. Intervals, where given, are 2.5th to 97.5th percentiles across "
  "replications rather than analytic approximations. Bootstrap intervals on the empirical results "
  "use 1,200 to 1,500 resamples of the pooled out-of-fold predictions.")

doc.save(OUT)
print("wrote", OUT)
